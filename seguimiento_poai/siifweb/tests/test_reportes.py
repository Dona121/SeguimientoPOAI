# -*- coding: utf-8 -*-
"""Los reportes: el indice y los tres reportes (Financiero, Plan indicativo y Proyectos
programados), cada uno con su pagina y su Excel.

La base de pruebas ya trae el catalogo de metas y la programacion: las migraciones 0012 y
0014 los cargan al crearla.
"""
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from siifweb import reportes
from siifweb.models import LineaEstrategica, MetaIndicador, ProgramacionMeta


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.usuario = get_user_model().objects.create_superuser("admin", "a@b.co", "clave")

    def setUp(self):
        self.client.force_login(self.usuario)

    def excel(self, nombre_url, **filtros):
        respuesta = self.client.get(reverse(nombre_url), {"descargar": "xlsx", **filtros})
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("attachment;", respuesta["Content-Disposition"])
        hoja = load_workbook(BytesIO(respuesta.content)).active
        encabezados = [c.value for c in hoja[2]]
        filas = [dict(zip(encabezados, [c.value for c in f])) for f in hoja.iter_rows(min_row=3)]
        return hoja, encabezados, filas


class ElIndice(Base):
    def test_trae_los_tres_reportes(self):
        respuesta = self.client.get(reverse("admin:siifweb_proyecto_reportes"))
        self.assertEqual(respuesta.status_code, 200)
        for nombre in ("reporte_financiero", "reporte_plan", "reporte_proyectos"):
            with self.subTest(reporte=nombre):
                self.assertContains(respuesta, reverse(f"admin:siifweb_proyecto_{nombre}"))

    def test_el_menu_apunta_al_indice(self):
        from django.conf import settings
        enlaces = [str(i["link"]) for g in settings.UNFOLD["SIDEBAR"]["navigation"] for i in g["items"]]
        self.assertIn(reverse("admin:siifweb_proyecto_reportes"), enlaces)

    def test_las_vigencias_del_filtro_no_se_repiten(self):
        for nombre in ("reporte_financiero", "reporte_proyectos"):
            with self.subTest(reporte=nombre):
                filtros = self.client.get(reverse(f"admin:siifweb_proyecto_{nombre}")).context["filtros"]
                vigencias = next(f for f in filtros if f["tipo"] == "vigencias")["opciones"]
                self.assertEqual(len(vigencias), len(set(vigencias)))
                self.assertEqual(vigencias, sorted(vigencias))

    def test_cada_reporte_tiene_su_pagina(self):
        for nombre in ("reporte_financiero", "reporte_plan", "reporte_proyectos"):
            with self.subTest(reporte=nombre):
                respuesta = self.client.get(reverse(f"admin:siifweb_proyecto_{nombre}"))
                self.assertEqual(respuesta.status_code, 200)
                self.assertContains(respuesta, 'name="descargar" value="xlsx"')
                self.assertContains(respuesta, reverse("admin:siifweb_proyecto_reportes"))


class LosPermisos(Base):
    def entrar_como(self, rol):
        usuario = get_user_model().objects.create_user(f"u_{rol}", "x@b.co", "clave", is_staff=True)
        if rol:
            usuario.groups.add(Group.objects.get(name=rol))
        self.client.force_login(usuario)

    def test_los_roles_solo_ven_el_financiero(self):
        for rol in ("Consulta", "Seguimiento"):
            with self.subTest(rol=rol):
                self.entrar_como(rol)
                respuesta = self.client.get(reverse("admin:siifweb_proyecto_reportes"))
                self.assertContains(respuesta, reverse("admin:siifweb_proyecto_reporte_financiero"))
                self.assertNotContains(respuesta, reverse("admin:siifweb_proyecto_reporte_plan"))
                self.assertEqual(self.client.get(reverse("admin:siifweb_proyecto_reporte_plan")).status_code, 403)
                self.assertEqual(self.client.get(reverse("admin:siifweb_proyecto_reporte_proyectos")).status_code, 403)

    def test_sin_ningun_reporte_el_indice_da_403(self):
        self.entrar_como(None)
        self.assertEqual(self.client.get(reverse("admin:siifweb_proyecto_reportes")).status_code, 403)


class PlanIndicativo(Base):
    URL = "admin:siifweb_proyecto_reporte_plan"

    def test_una_fila_por_meta_en_orden(self):
        hoja, encabezados, filas = self.excel(self.URL)
        self.assertEqual(len(filas), MetaIndicador.objects.count())
        codigos = [f["Código Meta"] for f in filas]
        self.assertEqual(codigos, sorted(codigos))
        self.assertEqual(encabezados, [c[1] for c in reportes.columnas_plan()])
        self.assertEqual(hoja.title, "Plan indicativo")

    def test_las_columnas_de_qrypi_que_hay_en_los_modelos(self):
        _, encabezados, _ = self.excel(self.URL)
        for columna in ("Línea Estratégica", "Indicador de resultado", "ODS", "Sector PDD",
                        "Número Programa PDD", "Programa PDD", "Sector de inversión",
                        "Programa Presupuestal", "Producto", "Indicador de producto principal",
                        "Meta de cuatrienio", "Responsable", "Meta Física Esperada 2024",
                        "Meta Física Esperada 2027"):
            with self.subTest(columna=columna):
                self.assertIn(columna, encabezados)

    def test_los_datos_de_una_meta(self):
        meta = MetaIndicador.objects.get(codigo="MT323")
        _, _, filas = self.excel(self.URL, mt="MT323")
        self.assertEqual(len(filas), 1)
        fila = filas[0]
        self.assertEqual(fila["Código del indicador principal"], meta.indicador_mga.codigo)
        self.assertEqual(Decimal(str(fila["Meta de cuatrienio"])), meta.meta_cuatrenio)
        for v in meta.vigencias.all():
            self.assertEqual(Decimal(str(fila[f"Meta Física Esperada {v.vigencia}"])), v.valor)

    def test_el_filtro_de_linea(self):
        linea = LineaEstrategica.objects.first()
        _, _, filas = self.excel(self.URL, linea=linea.pk)
        esperadas = MetaIndicador.objects.filter(programa_pdd__sector_pdd__linea_estrategica=linea).count()
        self.assertEqual(len(filas), esperadas)


class ProyectosProgramados(Base):
    URL = "admin:siifweb_proyecto_reporte_proyectos"

    def test_una_fila_por_programacion(self):
        hoja, encabezados, filas = self.excel(self.URL)
        self.assertEqual(len(filas), ProgramacionMeta.objects.count())
        self.assertEqual(encabezados, [c[1] for c in reportes.columnas_proyectos()])
        for columna in ("BPIN", "Vigencia PI", "Nombre del proyecto", "Tipo de banco",
                        "Meta del proyecto", "Municipio", "Código Meta", "Programa PDD"):
            with self.subTest(columna=columna):
                self.assertIn(columna, encabezados)

    def test_los_filtros_de_vigencia_y_mt(self):
        _, _, filas = self.excel(self.URL, vigencia="2025", mt="MT323")
        esperadas = ProgramacionMeta.objects.filter(vigencia=2025, meta__codigo="MT323").count()
        self.assertEqual(len(filas), esperadas)
        self.assertTrue(all(f["Vigencia PI"] == 2025 and f["Código Meta"] == "MT323" for f in filas))

    def test_la_meta_del_proyecto_va_exacta(self):
        _, _, filas = self.excel(self.URL, mt="MT323")
        self.assertIn(Decimal("3.0976"), [Decimal(str(f["Meta del proyecto"])) for f in filas
                                          if f["Meta del proyecto"] is not None])

    def test_no_hace_una_consulta_por_fila(self):
        """Con 1.609 filas, armar la meta de cada una por separado eran 3.226 consultas."""
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        with CaptureQueriesContext(connection) as consultas:
            self.excel(self.URL)
        self.assertLess(len(consultas), 30)
